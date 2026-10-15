"""
Entity-Based Sentence Generation.

Takes the rows produced by the existing NER step (app.py, section 03:
"Sentence No.", "Sentence", "Entity", "Entity Type", "Start", "End", ...)
and creates one new sentence per replacement: the entity at exactly that
character span is swapped for the user's replacement, everything else in
the sentence (text, punctuation, spacing, other occurrences of the same
words) stays as it was. Works on any Unicode text (English, Hindi,
Marathi). No NER is run here.

Exports the results as a Word document (python-docx, already a project
dependency) or plain text.
"""

import io
from typing import Any, Dict, List

ENTITY_TYPES = ("PER", "ORG", "LOC")


def replace_span(sentence: str, start: int, end: int, entity: str,
                 replacement: str) -> str:
    """
    `sentence` with sentence[start:end] replaced by `replacement`.
    Raises ValueError if the span does not hold `entity` (stale or wrong
    offsets), so a broken sentence is never produced.
    """
    if not (0 <= start < end <= len(sentence)) or sentence[start:end] != entity:
        raise ValueError(
            f"Entity '{entity}' is not at characters {start}-{end} of the sentence."
        )
    return sentence[:start] + replacement + sentence[end:]


def entity_key(row: Dict[str, Any]) -> str:
    """Stable identifier of one entity occurrence (sentence + span)."""
    return f"{row['Sentence No.']}:{row['Start']}:{row['End']}"


def generate_sentences(rows: List[Dict[str, Any]],
                       replacements: Dict[str, str]) -> List[Dict[str, Any]]:
    """
    One generated sentence per entity that has a non-empty replacement
    different from the original entity text. `replacements` maps
    entity_key(row) -> replacement text; missing or blank entries are
    skipped. Results are ordered by sentence and position.
    """
    generated = []

    for row in sorted(rows, key=lambda r: (r["Sentence No."], r["Start"])):
        replacement = (replacements.get(entity_key(row)) or "").strip()

        if not replacement or replacement == row["Entity"]:
            continue

        sentence = replace_span(row["Sentence"], row["Start"], row["End"],
                                row["Entity"], replacement)
        generated.append({
            "Sentence No.": row["Sentence No."],
            "Original Sentence": row["Sentence"],
            "Entity Type": row["Entity Type"],
            "Original Entity": row["Entity"],
            "Replacement": replacement,
            "Generated Sentence": sentence,
            # Span of the replacement inside the generated sentence.
            "Start": row["Start"],
            "End": row["Start"] + len(replacement),
        })

    return generated


def build_txt(generated: List[Dict[str, Any]]) -> str:
    """Plain-text export: each original sentence followed by its variants."""
    lines, current = ["Entity-Based Sentence Generation", ""], None

    for item in generated:
        if item["Original Sentence"] != current:
            current = item["Original Sentence"]
            if len(lines) > 2:
                lines.append("")
            lines.append(f"Original: {current}")
        lines.append(
            f"  {item['Entity Type']}: {item['Original Entity']} -> "
            f"{item['Replacement']}: {item['Generated Sentence']}"
        )

    return "\n".join(lines) + "\n"


def build_docx(generated: List[Dict[str, Any]]) -> bytes:
    """
    Word export: each original sentence, then a numbered list of its
    generated sentences with the replacement in bold and highlighted.
    """
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    document = Document()
    document.add_heading("Entity-Based Sentence Generation", level=1)
    current = None

    for item in generated:
        if item["Original Sentence"] != current:
            current = item["Original Sentence"]
            paragraph = document.add_paragraph()
            paragraph.add_run("Original: ").bold = True
            paragraph.add_run(current)

        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(
            f"{item['Entity Type']}: {item['Original Entity']} → "
            f"{item['Replacement']} — "
        ).italic = True
        sentence, start, end = item["Generated Sentence"], item["Start"], item["End"]
        paragraph.add_run(sentence[:start])
        replaced = paragraph.add_run(sentence[start:end])
        replaced.bold = True
        replaced.font.highlight_color = WD_COLOR_INDEX.YELLOW
        paragraph.add_run(sentence[end:])

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
