"""
Entity-Based Sentence Generation tests (sentence_generation.py and the
"Generate New Sentences" section of app.py).

Run from the project folder with the Python 3.12 environment:

    .venv312\\Scripts\\python tests\\test_generation.py

Part 1 (deterministic): replacement, skipping and exports, with NER rows
written by hand in the format of app.py's NER results.
Part 2: the real Streamlit app (NER results from the existing pipeline).
Exit code is 0 only if every check passes.
"""

import io
import os
import sys
import logging
import warnings

warnings.filterwarnings("ignore")
logging.disable(logging.WARNING)

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(HERE)
APP_PATH = os.path.join(PROJECT, "app.py")

sys.path.insert(0, PROJECT)
sys.stdout.reconfigure(encoding="utf-8")

from docx import Document  # noqa: E402
from docx.enum.text import WD_COLOR_INDEX  # noqa: E402

from sentence_generation import (  # noqa: E402
    build_docx, build_txt, entity_key, generate_sentences, replace_span,
)

failures = []
passes = 0


def check(name, condition, detail=""):
    global passes
    if condition:
        passes += 1
        print(f"  PASS  {name}")
    else:
        failures.append(name)
        print(f"  FAIL  {name}\n        -> {detail}")


def section(title):
    print(f"\n=== {title}")


def rows_for(sentence, entities, number=1):
    """NER rows (app.py format) for (text, type) entities, found left to right."""
    rows, start = [], 0
    for text, entity_type in entities:
        s = sentence.index(text, start)
        rows.append({"Sentence No.": number, "Sentence": sentence, "Entity": text,
                     "Entity Type": entity_type, "Confidence": 0.9,
                     "Start": s, "End": s + len(text), "Source": "test"})
        start = s + len(text)
    return rows


def by_entity(rows, mapping):
    """Replacement dict keyed like the app: entity_key(row) -> text."""
    return {entity_key(r): mapping[r["Entity"]] for r in rows if r["Entity"] in mapping}


# ------------------------------------------------------------------
section("Part 1: replacement logic")

EXAMPLE = "Aarav Sharma works in Infosys in Pune."
rows = rows_for(EXAMPLE, [("Aarav Sharma", "PER"), ("Infosys", "ORG"), ("Pune", "LOC")])
generated = generate_sentences(rows, by_entity(rows, {
    "Aarav Sharma": "Manya", "Infosys": "Microsoft", "Pune": "Delhi"}))
check("example: one new sentence per replacement, only that entity changed",
      [g["Generated Sentence"] for g in generated] == [
          "Manya works in Infosys in Pune.",
          "Aarav Sharma works in Microsoft in Pune.",
          "Aarav Sharma works in Infosys in Delhi."],
      [g["Generated Sentence"] for g in generated])
check("each result records type, original entity and replacement",
      [(g["Entity Type"], g["Original Entity"], g["Replacement"]) for g in generated]
      == [("PER", "Aarav Sharma", "Manya"), ("ORG", "Infosys", "Microsoft"),
          ("LOC", "Pune", "Delhi")])
check("Start/End point at the replacement in the generated sentence",
      all(g["Generated Sentence"][g["Start"]:g["End"]] == g["Replacement"] for g in generated))

sentence = "Pune is near Mumbai, and Pune  is green."
rows2 = rows_for(sentence, [("Pune", "LOC"), ("Mumbai", "LOC"), ("Pune", "LOC")])
second_pune = {entity_key(rows2[2]): "Nashik"}
check("same text twice: only the intended occurrence is replaced; spacing kept",
      generate_sentences(rows2, second_pune)[0]["Generated Sentence"]
      == "Pune is near Mumbai, and Nashik  is green.")

HI = "विक्रम पटेल पुणे में प्रिया मेहता से मिले।"
rows_hi = rows_for(HI, [("विक्रम पटेल", "PER"), ("पुणे", "LOC"), ("प्रिया मेहता", "PER")])
out = generate_sentences(rows_hi, by_entity(rows_hi, {"पुणे": "दिल्ली", "प्रिया मेहता": "मान्या"}))
check("Hindi (Devanagari): replacements and punctuation preserved",
      [g["Generated Sentence"] for g in out] == [
          "विक्रम पटेल दिल्ली में प्रिया मेहता से मिले।",
          "विक्रम पटेल पुणे में मान्या से मिले।"], [g["Generated Sentence"] for g in out])

MR = "आरव शर्मा पुण्यात इन्फोसिसमध्ये काम करतात।"
rows_mr = rows_for(MR, [("आरव शर्मा", "PER"), ("इन्फोसिस", "ORG")])
out = generate_sentences(rows_mr, by_entity(rows_mr, {"इन्फोसिस": "मायक्रोसॉफ्ट"}))
check("Marathi: entity inside an inflected word replaced, case suffix kept",
      [g["Generated Sentence"] for g in out] == ["आरव शर्मा पुण्यात मायक्रोसॉफ्टमध्ये काम करतात।"],
      [g["Generated Sentence"] for g in out])

check("blank, whitespace-only and missing replacements are skipped",
      generate_sentences(rows, {entity_key(rows[0]): "", entity_key(rows[1]): "   "}) == [])
check("a replacement equal to the original entity is skipped",
      generate_sentences(rows, {entity_key(rows[0]): "Aarav Sharma"}) == [])
check("replacement text is trimmed",
      generate_sentences(rows, {entity_key(rows[2]): "  Delhi "})[0]["Generated Sentence"]
      == "Aarav Sharma works in Infosys in Delhi.")
only_per = rows_for("Manya lives here.", [("Manya", "PER")])
check("sentence without ORG/LOC entities works (missing types are fine)",
      generate_sentences(only_per, {entity_key(only_per[0]): "Ravi"})[0]["Generated Sentence"]
      == "Ravi lives here.")
check("no NER rows -> nothing generated", generate_sentences([], {}) == [])

multi = rows_for("Ravi met Asha.", [("Asha", "PER")], number=2) + rows
out = generate_sentences(multi, by_entity(multi, {"Asha": "Meera", "Pune": "Delhi"}))
check("several sentences: results ordered by sentence number",
      [g["Sentence No."] for g in out] == [1, 2], [g["Sentence No."] for g in out])

stale = dict(rows[0], Start=1)
try:
    generate_sentences([stale], {entity_key(stale): "Manya"})
    raised = False
except ValueError:
    raised = True
check("offsets that do not match the entity are rejected (no broken sentence)", raised)
check("replace_span replaces exactly the given span",
      replace_span("A Pune B Pune", 9, 13, "Pune", "Delhi") == "A Pune B Delhi")

# ------------------------------------------------------------------
section("Part 1: exports")

mixed = generated + generate_sentences(rows_hi, by_entity(rows_hi, {"पुणे": "दिल्ली"}))
doc = Document(io.BytesIO(build_docx(mixed)))
texts = [p.text for p in doc.paragraphs]
check("Word: heading, original sentences and every generated sentence",
      texts[0] == "Entity-Based Sentence Generation"
      and f"Original: {EXAMPLE}" in texts and f"Original: {HI}" in texts
      and all(any(g["Generated Sentence"] in t for t in texts) for g in mixed), texts)
highlights = [r.text for p in doc.paragraphs for r in p.runs
              if r.bold and r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
check("Word: each replacement is the bold, highlighted run (Unicode kept)",
      highlights == [g["Replacement"] for g in mixed], highlights)
check("Word: generated sentences use a numbered list",
      sum(p.style.name == "List Number" for p in doc.paragraphs) == len(mixed))

txt = build_txt(mixed)
check("TXT: originals and generated sentences, UTF-8 safe",
      f"Original: {EXAMPLE}" in txt and "विक्रम पटेल दिल्ली में प्रिया मेहता से मिले।" in txt
      and "PER: Aarav Sharma -> Manya: Manya works in Infosys in Pune." in txt
      and txt.encode("utf-8").decode("utf-8") == txt, txt)

# ------------------------------------------------------------------
section("Part 2: the real app (Monolingual NER Extraction)")

from streamlit.testing.v1 import AppTest  # noqa: E402

REPLACE = {"Aarav Sharma": "Manya", "Infosys": "Microsoft", "Pune": "Delhi"}


def extract(text, language):
    at = AppTest.from_file(APP_PATH, default_timeout=900)
    at.run()
    at.selectbox(key="ner_language").set_value(language).run()
    at.text_area(key="manual_text").set_value(text)
    at.button(key="process_input").click().run()
    at.button(key="run_ner").click().run()
    return at


def ner_rows(at):
    table = next((d.value for d in at.dataframe if "Entity Type" in d.value.columns
                  and "Source" in d.value.columns), None)
    return [] if table is None else table.to_dict("records")


at = extract(EXAMPLE, "English")
found = ner_rows(at)
print("  NER rows:", [(r["Entity"], r["Entity Type"]) for r in found])
check("section 'Generate New Sentences' follows the NER results",
      "Generate New Sentences" in [h.value for h in at.header]
      and any(b.key == "generate_sentences" for b in at.button))
fields = {t.key: t for t in at.text_input if t.key and t.key.startswith("gen_")}
check("one replacement field per extracted entity, labelled with the entity",
      len(fields) == len(found)
      and all(any(f"“{r['Entity']}”" in t.label for t in fields.values()) for r in found),
      [t.label for t in fields.values()])
groups = [m.value for m in at.markdown if m.value.startswith("**") and " · " in m.value]
check("entities grouped by type in PER, ORG, LOC order",
      groups == [g for g in ("**PER** · person", "**ORG** · organization", "**LOC** · location")
                 if g.split("**")[1] in {r["Entity Type"] for r in found}], groups)

at.button(key="generate_sentences").click().run()
check("no replacements entered -> clear warning, nothing generated",
      any("Enter at least one replacement" in w.value for w in at.warning)
      and not any(s.value.startswith("Generated") for s in at.success))

# Widget handles are per run: collect the fields again after the click.
fields = {t.key: t for t in at.text_input if t.key and t.key.startswith("gen_")}
for key, field in fields.items():
    entity = key.rsplit("_", 1)[1]
    if entity in REPLACE:
        field.set_value(REPLACE[entity])
at.button(key="generate_sentences").click().run()
table = next((d.value for d in at.dataframe if "Generated Sentence" in d.value.columns), None)
got = [] if table is None else list(table["Generated Sentence"])
expected = [EXAMPLE.replace(e, REPLACE[e]) for e in REPLACE
            if any(r["Entity"] == e for r in found)]
check("app: one new sentence per replaced entity, nothing else changed",
      not at.exception and sorted(got) == sorted(expected) and len(got) == len(found),
      (got, expected))
check("app: success message and original sentence shown",
      any(s.value == f"Generated {len(got)} new sentences." for s in at.success)
      and any(m.value == f"**Original:** {EXAMPLE.replace('.', chr(92) + '.')}"
              for m in at.markdown))
check("app: replaced entity highlighted in the result",
      any(":orange-background[Manya]" in m.value for m in at.markdown))
check("app: Word and TXT downloads offered",
      {b.label for b in at.get("download_button")}
      >= {"Download Sentences (Word)", "Download Sentences (TXT)"})

at.button(key="run_ner").click().run()
check("a new extraction clears earlier generated sentences",
      not any("Generated Sentence" in d.value.columns for d in at.dataframe))

at = extract(HI, "Hindi")
found = ner_rows(at)
fields = {t.key: t for t in at.text_input if t.key and t.key.startswith("gen_")}
pune = next((k for k in fields if k.endswith("_पुणे")), None)
check("Hindi: पुणे offered for replacement", pune is not None, list(fields))
if pune:
    fields[pune].set_value("दिल्ली")
    at.button(key="generate_sentences").click().run()
    table = next((d.value for d in at.dataframe if "Generated Sentence" in d.value.columns), None)
    check("Hindi: only पुणे replaced in the generated sentence",
          table is not None
          and list(table["Generated Sentence"]) == ["विक्रम पटेल दिल्ली में प्रिया मेहता से मिले।"],
          None if table is None else list(table["Generated Sentence"]))

print(f"\nRESULT: {passes} passed, {len(failures)} failed")
for name in failures:
    print("  failed:", name)
sys.exit(1 if failures else 0)
